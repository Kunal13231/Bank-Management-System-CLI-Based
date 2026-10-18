import openpyxl

file = "bank.xlsx"

# Create Excel file
try:
    wkb = openpyxl.load_workbook(file)
    ws = wkb.active
except FileNotFoundError:
    wkb = openpyxl.Workbook()
    ws = wkb.active
    ws.append(["Name", "Mobile", "Balance"])
    wkb.save(file)
    


def add_new():
    wkb = openpyxl.load_workbook(file)
    ws = wkb.active

    name = input("Enter Your Name: ")
    mobile = input("Enter Mobile Number: ")
    balance = float(input("Enter Opening Balance: "))

    ws.append([name, mobile, balance])
    wkb.save(file)
    wkb.close()

    print("Account Created Successfully!")


def deposit_money():
    wkb = openpyxl.load_workbook(file)
    ws = wkb.active

    mobile = input("Enter Mobile Number: ")
    amount = float(input("Enter Amount to Deposit: "))

    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row, 2).value) == mobile:
            ws.cell(row, 3).value += amount
            wkb.save(file)
            wkb.close()
            print("Money Deposited Successfully!")
            return

    print("Account Not Found!")
    wkb.close()


def withdrawal_money():
    wkb = openpyxl.load_workbook(file)
    ws = wkb.active

    mobile = input("Enter Mobile Number: ")
    amount = float(input("Enter Amount to Withdraw: "))

    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row, 2).value) == mobile:
            balance = ws.cell(row, 3).value

            if amount <= 0:
                print("Enter a Valid Amount!")
            elif amount <= balance:
                ws.cell(row, 3).value = balance - amount
                wkb.save(file)
                print("Money Withdrawn Successfully!")
            else:
                print("Insufficient Balance!")

            wkb.close()
            return

    print("Account Not Found!")
    wkb.close()


def view_balance():
    wkb = openpyxl.load_workbook(file)
    ws = wkb.active

    mobile = input("Enter Mobile Number: ")

    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row, 2).value) == mobile:
            print("Name:", ws.cell(row, 1).value)
            print("Current Balance: Rs.", ws.cell(row, 3).value)
            wkb.close()
            return

    print("Account Not Found!")
    wkb.close()


def menu():
    while True:
        print("\nWELCOME TO THE BANK..!")
        print("1. Add New Account")
        print("2. Deposit Money")
        print("3. Withdraw Money")
        print("4. View Balance")
        print("5. Exit")

        choice = input("Enter Your Choice: ")

        if choice == "1":
            add_new()
        elif choice == "2":
            deposit_money()
        elif choice == "3":
            withdrawal_money()
        elif choice == "4":
            view_balance()
        elif choice == "5":
            print("Thank You!")
            print("Rate My Code...!")
            break
        else:
            print("Invalid Choice!")


menu()

